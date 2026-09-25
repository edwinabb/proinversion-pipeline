"""Genera el Excel de pipeline y el reporte Markdown de cambios."""

from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_DIR = Path("output")
PROCESSED_DIR = Path("data/processed")

PRIORITY_COLORS = {
    "A": "FF2E75B6",
    "B": "FF70AD47",
    "C": "FFFFC000",
}

ACTIONS = {
    "TEMPRANO": "Solicitar reunión de sondeo de mercado con la entidad",
    "MEDIO": "Preparar propuesta técnica y gestionar reunión con ProInversión",
    "URGENTE": "Revisar TDR, verificar requisitos y preparar propuesta para proceso de selección",
    "COMPETENCIA": "Mapear empresa adjudicataria; evaluar participación como subcontratista",
}


def _apply_table_style(ws, n_rows: int, n_cols: int):
    """Congela encabezado, aplica filtros automáticos y formato."""
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{n_rows + 1}"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF203864")
        cell.alignment = Alignment(wrap_text=True)
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, 40)


def _write_df_to_sheet(ws, df: pd.DataFrame, url_col: str = "fuente_url"):
    ws.append(list(df.columns))
    for _, row in df.iterrows():
        ws.append([str(v) if pd.notna(v) else "" for v in row])
    # Hipervínculos en columna URL
    if url_col in df.columns:
        col_idx = list(df.columns).index(url_col) + 1
        for i, val in enumerate(df[url_col], start=2):
            if isinstance(val, str) and val.startswith("http"):
                cell = ws.cell(row=i, column=col_idx)
                cell.hyperlink = val
                cell.font = Font(color="FF0563C1", underline="single")
    _apply_table_style(ws, len(df), len(df.columns))


def generate_excel(df: pd.DataFrame, run_date: date, seed_results: list, manifests: dict) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out_path = OUTPUT_DIR / f"pipeline_{run_date.isoformat()}.xlsx"

    display_cols = [
        "nombre", "modalidad", "sector_funcion", "entidad", "region",
        "monto_usd", "monto_soles", "fase_estado", "timing", "fit_nivel",
        "fit_razon", "fecha_adjudicacion_estimada", "cui", "fuente_url",
        "campos_faltantes", "portal",
    ]

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        for prio in ("A", "B", "C"):
            subset = df[df["prioridad"] == prio].sort_values("monto_usd", ascending=False, na_position="last")
            cols = [c for c in display_cols if c in subset.columns]
            subset[cols].to_excel(writer, sheet_name=f"Prioridad_{prio}", index=False)

        # Competencia: timing == COMPETENCIA
        comp = df[df["timing"] == "COMPETENCIA"].sort_values("monto_usd", ascending=False, na_position="last")
        cols = [c for c in display_cols if c in comp.columns]
        comp[cols].to_excel(writer, sheet_name="Competencia", index=False)

        # Resumen
        resumen_data = {
            "Por modalidad": df[df["prioridad"].notna()].groupby("modalidad").size().reset_index(name="n"),
            "Por región": df[df["prioridad"].notna()].groupby("region").size().reset_index(name="n"),
            "Por prioridad": df[df["prioridad"].notna()].groupby("prioridad").size().reset_index(name="n"),
            "Por timing": df[df["prioridad"].notna()].groupby("timing").size().reset_index(name="n"),
        }
        row = 1
        resumen_ws = writer.book.create_sheet("Resumen")
        for title, sub_df in resumen_data.items():
            resumen_ws.cell(row=row, column=1, value=title).font = Font(bold=True)
            row += 1
            resumen_ws.append(list(sub_df.columns))
            for _, r in sub_df.iterrows():
                resumen_ws.append(list(r))
                row += 1
            row += 1

        # Calidad de datos
        calidad = df[df["campos_faltantes"].str.len() > 0][
            ["nombre", "portal", "campos_faltantes", "fuente_url"]
        ] if "campos_faltantes" in df.columns else pd.DataFrame()
        calidad.to_excel(writer, sheet_name="Calidad_datos", index=False)

    # Aplicar estilos post-escritura
    wb = load_workbook(out_path)
    for prio in ("A", "B", "C"):
        sname = f"Prioridad_{prio}"
        if sname in wb.sheetnames:
            ws = wb[sname]
            _apply_table_style(ws, ws.max_row - 1, ws.max_column)
    wb.save(out_path)

    print(f"Excel generado: {out_path}")
    return out_path


def compute_diff(df_current: pd.DataFrame, run_date: date) -> dict:
    hist_file = PROCESSED_DIR / "oportunidades.parquet"
    if not hist_file.exists():
        return {"nuevos": [], "cambios_fase": [], "desaparecidos": []}

    hist = pd.read_parquet(hist_file)
    prev_dates = sorted(hist["fecha_extraccion"].unique())
    if len(prev_dates) < 2:
        return {"nuevos": list(df_current["nombre"].head(5)), "cambios_fase": [], "desaparecidos": []}

    prev_date = prev_dates[-2]
    df_prev = hist[hist["fecha_extraccion"] == prev_date]

    current_names = set(df_current["nombre"].str.lower())
    prev_names = set(df_prev["nombre"].str.lower())

    nuevos = list(df_current[~df_current["nombre"].str.lower().isin(prev_names)]["nombre"].head(10))
    desaparecidos = list(df_prev[~df_prev["nombre"].str.lower().isin(current_names)]["nombre"].head(10))

    # Cambios de fase
    cambios = []
    for _, row in df_current.iterrows():
        match = df_prev[df_prev["nombre"].str.lower() == str(row["nombre"]).lower()]
        if not match.empty:
            prev_fase = match.iloc[0].get("fase_estado", "")
            curr_fase = row.get("fase_estado", "")
            if prev_fase != curr_fase:
                cambios.append({"nombre": row["nombre"], "fase_anterior": prev_fase, "fase_nueva": curr_fase})

    return {"nuevos": nuevos, "cambios_fase": cambios[:10], "desaparecidos": desaparecidos}


def generate_report(
    df: pd.DataFrame,
    run_date: date,
    seed_results: list,
    manifests: dict,
    dubious: list,
) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    diff = compute_diff(df, run_date)

    top_a = df[df["prioridad"] == "A"].sort_values("monto_usd", ascending=False, na_position="last").head(5)
    seeds_missing = [s for s in seed_results if not s["encontrada"]]

    errores = []
    for pid, m in manifests.items():
        if m.get("errores"):
            errores.append(f"- **{pid}**: {'; '.join(m['errores'])}")
        if m.get("total_ui") and m.get("total_extraido") and m["total_ui"] != m["total_extraido"]:
            errores.append(f"- **{pid}**: UI reporta {m['total_ui']} registros, extraídos {m['total_extraido']}")

    lines = [
        f"# Reporte de pipeline ProInversión — {run_date.isoformat()}\n",
        "## Resumen ejecutivo\n",
        f"- Total registros con fit ≥ 1: {(df['prioridad'].notna()).sum()}",
        f"- Prioridad A: {(df['prioridad'] == 'A').sum()} | B: {(df['prioridad'] == 'B').sum()} | C: {(df['prioridad'] == 'C').sum()}",
        f"- Total > US$ 5M: {(df['monto_usd'] > 5_000_000).sum() if 'monto_usd' in df.columns else 'N/D'}",
        "",
        "## Novedades frente a la corrida anterior\n",
        "**Nuevos:** " + (", ".join(diff["nuevos"][:5]) if diff["nuevos"] else "Ninguno"),
        "",
        "**Cambios de fase:**",
    ]

    if diff["cambios_fase"]:
        for c in diff["cambios_fase"]:
            lines.append(f"- {c['nombre'][:50]}: {c['fase_anterior']} → {c['fase_nueva']}")
    else:
        lines.append("- Ninguno")

    lines += [
        "",
        "**Desaparecidos:** " + (", ".join(diff["desaparecidos"][:5]) if diff["desaparecidos"] else "Ninguno"),
        "",
        "## Top 5 Prioridad A\n",
    ]

    for i, (_, row) in enumerate(top_a.iterrows(), 1):
        monto = f"US$ {row['monto_usd']:,.0f}" if pd.notna(row.get("monto_usd")) else "monto desconocido"
        accion = ACTIONS.get(row.get("timing", ""), "Evaluar oportunidad")
        fecha = row.get("fecha_adjudicacion_estimada", "N/D")
        lines += [
            f"### {i}. {row['nombre']}",
            f"- **Monto:** {monto} | **Fase:** {row.get('fase_estado', 'N/D')} | **Timing:** {row.get('timing', 'N/D')}",
            f"- **Entidad:** {row.get('entidad', 'N/D')} | **Región:** {row.get('region', 'N/D')}",
            f"- **Acción:** {accion}",
            f"- **Fecha límite estimada:** {fecha}",
            f"- **Fuente:** {row.get('fuente_url', 'N/D')}",
            "",
        ]

    lines += ["## Semillas no encontradas\n"]
    if seeds_missing:
        for s in seeds_missing:
            lines.append(f"- {s['semilla']}")
    else:
        lines.append("- Todas las semillas encontradas.")

    lines += ["", "## Problemas\n"]
    if errores:
        lines.extend(errores)
    else:
        lines.append("- Sin problemas de extracción reportados.")

    if dubious:
        lines += ["", "## Casos dudosos para revisión humana\n"]
        for d in dubious[:10]:
            lines.append(f"- {d['nombre'][:60]} | {d.get('sector', '')} | {d.get('fit_razon', '')}")

    report_path = OUTPUT_DIR / f"reporte_{run_date.isoformat()}.md"
    report_path.write_text("\n".join(lines))
    print(f"Reporte generado: {report_path}")
    return report_path
